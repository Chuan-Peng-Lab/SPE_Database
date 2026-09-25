# -*- coding: utf-8 -*-
"""Consistency pass over SPE_数据库_v16.2.docx."""
import os
import re

import docx

DST = r'D:\GitHub_programe\GitHub\SPE_Database\Datasets\4_Writing\SPE_数据库_v16.2.docx'
PIC = r'D:\GitHub_programe\GitHub\SPE_Database\3_Reports\Output\Pic'
d = docx.Document(DST)

print("== structure ==")
print("paragraphs:", len(d.paragraphs), " tables:", len(d.tables),
      " inline shapes:", len(d.inline_shapes))
for i, t in enumerate(d.tables):
    print(f"  table {i}: {len(t.rows)}x{len(t.columns)}")

full = "\n".join(p.text for p in d.paragraphs)

print("\n== stale-value scan (should all be 0) ==")
STALE = ["44 studies", "70 experiments", "3,603", "1,554,083", "1.55 million",
         "0.298", "0.126", "0.066", "0.061", "0.182", "0.023", "0.028", "0.013",
         "8 of the 10", "\u2212.493", ".374", ".389", "table 3", "Ratchliff",
         "Rose, H.", "4 studies with", "1, 676", "853,015", "303, 331", "199,271",
         "Paper_Id", "_raw_Subject.csv", "Exp1_Clean.xlsx", "five categories"]
for s in STALE:
    n = full.count(s)
    flag = "OK " if n == 0 else "!! "
    if n:
        print(f"  {flag}{s!r}: {n}")
print("  (nothing listed above = clean)")

print("\n== cross-reference inventory ==")
for m in set(re.findall(r"(Figure\s*\d+[A-D]?|Figures\s*\d+[A-D]?\s*and\s*\d+[A-D]?|Table\s*\d+)", full)):
    print("  ", m, full.count(m))

print("\n== Figure captions ==")
for p in d.paragraphs:
    if p.style.name == "Caption" or p.text.strip().startswith("Figure "):
        print("  *", p.text.strip()[:130])

print("\n== Table 1 head / tail ==")
t1 = d.tables[0]
for r in [0, 1, 2, len(t1.rows) - 2, len(t1.rows) - 1]:
    print("  ", " | ".join(c.text.strip()[:26] for c in t1.rows[r].cells))

print("\n== Table 2 ==")
t2 = d.tables[1]
for r in t2.rows:
    cells, seen = [], set()
    for c in r.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        cells.append(c.text.strip())
    print("  ", " | ".join(cells))

print("\n== reference list ==")
refs = [p.text.strip() for p in d.paragraphs if p.style.name.startswith("Bibliography") and p.text.strip()]
print("count:", len(refs))
for r in refs:
    print("  -", r[:120])

print("\n== embedded images ==")
for i, s in enumerate(d.inline_shapes):
    part = d.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed]
    print(f"  {i}: {part.partname} {len(part.blob)} bytes  display "
          f"{s.width}x{s.height}")

print("\n== example-section numbers present ==")
for key in ["49 studies", "89 experiment-level", "4,875", "2,131,108",
            "1,068,081", "4,520", "0.252", "0.123", "0.052", "0.049",
            "0.125", "0.005", "0.009", "371,894", "276,010",
            "\u2212.311", "\u2212.335", ".045", ".174"]:
    print(f"  {key!r}: {full.count(key)}")
