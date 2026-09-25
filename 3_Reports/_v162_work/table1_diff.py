# -*- coding: utf-8 -*-
"""Compare the manuscript v16.1 Table 1 (in the docx) with the regenerated Table1_v2.csv."""
import csv
import io
import os
import re

import docx
import pandas as pd
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOCX = r'D:\GitHub_programe\GitHub\SPE_Database\Datasets\4_Writing\SPE_数据库_v16.1.docx'
NEW = r'D:\GitHub_programe\GitHub\SPE_Database\3_Reports\Output\Table1_v2.csv'
OUT = r'D:\GitHub_programe\GitHub\SPE_Database\3_Reports\_v162_work\table1_diff.txt'


def cell_text(cell):
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


d = docx.Document(DOCX)
t = d.tables[0]
old = [[cell_text(c) for c in row.cells] for row in t.rows]
old_header = old[0]
old_rows = old[1:]
print("manuscript Table 1:", len(old_rows), "data rows; header:", old_header)

new = list(csv.DictReader(open(NEW, encoding="utf-8-sig")))
print("regenerated Table 1:", len(new), "data rows")

# --- overall set differences ----------------------------------------------
def key(r, idcol, expcol):
    return (r[idcol], r[expcol])

old_keys = [(r[1], r[3]) for r in old_rows]
new_keys = [(r["ID"], r["Exp"]) for r in new]

from collections import Counter
oc, nc = Counter(old_keys), Counter(new_keys)

buf = []
buf.append(f"manuscript Table 1 : {len(old_rows)} rows")
buf.append(f"regenerated Table 1: {len(new)} rows")
buf.append("")
buf.append("== studies only in the regenerated table ==")
for k in sorted(set(new_keys) - set(old_keys)):
    buf.append(f"  + {k[0]} {k[1]}")
buf.append("")
buf.append("== studies only in the manuscript table ==")
for k in sorted(set(old_keys) - set(new_keys)):
    buf.append(f"  - {k[0]} {k[1]}")
buf.append("")
buf.append("== row-count differences per (ID, Exp) ==")
for k in sorted(set(oc) | set(nc)):
    if oc.get(k, 0) != nc.get(k, 0):
        buf.append(f"  {k[0]} {k[1]}: manu {oc.get(k,0)} -> new {nc.get(k,0)}")

# --- cell-level comparison on rows present in both (matched in order) -----
buf.append("")
buf.append("== cell-level differences (rows matched by (ID, Exp), in order) ==")
old_by = {}
for r in old_rows:
    old_by.setdefault((r[1], r[3]), []).append(r)
new_by = {}
for r in new:
    new_by.setdefault((r["ID"], r["Exp"]), []).append(r)

n_diff = 0
for k in sorted(set(old_by) & set(new_by)):
    o_list, n_list = old_by[k], new_by[k]
    for i in range(min(len(o_list), len(n_list))):
        o, n = o_list[i], n_list[i]
        for ci, col in [(5, "Country"), (6, "Language"), (7, "N (M/F)"),
                        (8, "Stimulus"), (9, "Trials"), (10, "License"),
                        (11, "Exp_Implement")]:
            ov = o[ci - 1]
            nv = str(n[col]) if n[col] != "nan" else ""
            if ov.strip() != nv.strip():
                n_diff += 1
                buf.append(f"  {k[0]} {k[1]} [{col}] manu='{ov}' -> new='{nv}'")
buf.append(f"\ntotal cell differences: {n_diff}")

# --- per-column mismatch summary -----------------------------------------
summary = Counter()
for line in buf:
    m = re.search(r"\[(.+?)\] manu=", line)
    if m:
        summary[m.group(1)] += 1
buf.append("\nby column: " + ", ".join(f"{k}={v}" for k, v in summary.most_common()))

open(OUT, "w", encoding="utf-8").write("\n".join(buf))
print("\n".join(buf[:80]))
print(f"\n... full report -> {OUT}")
