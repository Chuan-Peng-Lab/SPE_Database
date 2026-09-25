# -*- coding: utf-8 -*-
"""Batch 5: rebuild the caption paragraphs, removing Word's leftover SEQ-field runs
(which produced e.g. 'Figure 4 ... 4' and 'Table 23' after the text edit)."""
import copy

import docx
from docx.oxml.ns import qn

DST = r'D:\GitHub_programe\GitHub\SPE_Database\Datasets\4_Writing\SPE_数据库_v16.2.docx'
doc = docx.Document(DST)


def rebuild(p, text):
    """Replace the paragraph body with a single run that keeps the first run's formatting."""
    rPr = None
    runs = p.runs
    if runs:
        src = runs[0]._r.find(qn('w:rPr'))
        if src is not None:
            rPr = copy.deepcopy(src)
    for child in list(p._p):
        if child.tag != qn('w:pPr'):
            p._p.remove(child)
    r = p.add_run(text)
    if rPr is not None:
        r._r.insert(0, rPr)
    return p


CAPTIONS = {
    20: "Figure 1. Database folder structure.",
    37: "Table 1",
    45: "Figure 3. Standardized codebook of the experimental dataset variables.",
    51: ("Figure 4. Self-prioritization effects along the social distance continuum and pairwise "
         "comparisons across social identities."),
    64: ("Figure 5. Bootstrap estimation of the self-prioritization effect under nonmatching "
         "conditions. A and B, the conservative approach; C and D, the liberal approach."),
    65: "Table 2",
    78: ("Figure 6. Exploratory moderators of the self-prioritization effect under matching "
         "conditions."),
}
for i, t in CAPTIONS.items():
    rebuild(doc.paragraphs[i], t)
    print(f'[{i}] rebuilt')

doc.save(DST)


def visible(p):
    return ''.join((t.text or '') for t in p._p.iter(qn('w:t')))


d = docx.Document(DST)
print('\n== captions after fix ==')
for i in [20, 28, 37, 45, 51, 64, 65, 78]:
    print(f'  [{i}] {visible(d.paragraphs[i])!r}')

print('\n== paragraphs where visible != .text ==')
for i, p in enumerate(d.paragraphs):
    if p.text.strip() != visible(p).strip():
        print(f'  [{i}] .text={p.text[:90]!r}')
        print(f'        vis ={visible(p)[:110]!r}')
