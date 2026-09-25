# -*- coding: utf-8 -*-
"""Inventory the docx body: paragraph index, style, run count, text (truncated)."""
import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOCX = r'D:\GitHub_programe\GitHub\SPE_Database\Datasets\4_Writing\SPE_数据库_v16.1.docx'
OUT = r'D:\GitHub_programe\GitHub\SPE_Database\3_Reports\_v162_work\docx_inventory.txt'

d = docx.Document(DOCX)
lines = []
for i, p in enumerate(d.paragraphs):
    has_img = 'graphicData' in p._p.xml
    lines.append(f"[{i:3d}] style={p.style.name!r:28s} runs={len(p.runs):3d} img={has_img} "
                 f"| {p.text[:150]}")
lines.append("")
lines.append("== tables ==")
for ti, t in enumerate(d.tables):
    lines.append(f"table {ti}: {len(t.rows)}x{len(t.columns)} style={t.style.name if t.style else None}")
    lines.append("   row0: " + " | ".join(c.text.strip()[:24] for c in t.rows[0].cells))
    lines.append("   row1: " + " | ".join(c.text.strip()[:24] for c in t.rows[1].cells))
    lines.append("   last: " + " | ".join(c.text.strip()[:24] for c in t.rows[-1].cells))

txt = "\n".join(lines)
open(OUT, "w", encoding="utf-8").write(txt)
print(txt)
