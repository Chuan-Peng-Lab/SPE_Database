# -*- coding: utf-8 -*-
"""Batch 4: repair the Sun et al. (2023) entry and put the changed references in alphabetical order."""
import docx

DST = r'D:\GitHub_programe\GitHub\SPE_Database\Datasets\4_Writing\SPE_数据库_v16.2.docx'
doc = docx.Document(DST)


def set_text(p, text):
    if not p.runs:
        p.add_run(text); return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


biblio = [p for p in doc.paragraphs
          if p.style.name.startswith("Bibliography") and p.text.strip()]

# --- 1) the surviving Sun entry currently holds the wrong title ------------
sun = [p for p in biblio if p.text.strip().startswith("Sun, S.")]
print("Sun entries:", len(sun))
for p in sun:
    print("   before:", p.text[:90])
correct_sun = ("Sun, S., Wang, N., Wen, J., & Hu, C.-P. (2023). A dataset of cognitive ontology "
               "for neuroimaging studies of self-reference. China Scientific Data, 8(3), "
               "175\u2013189. https://doi.org/10.11922/11-6035.csd.2022.0047.zh")
set_text(sun[0], correct_sun)
for extra in sun[1:]:
    extra._element.getparent().remove(extra._element)
print("   after :", sun[0].text[:90])

# --- 2) alphabetical placement for the two changed / inserted entries ------
biblio = [p for p in doc.paragraphs
          if p.style.name.startswith("Bibliography") and p.text.strip()]
by_key = {}
for p in biblio:
    by_key.setdefault(p.text.strip()[:14], p)

aron = [p for p in biblio if p.text.startswith("Aron, A.")][0]
bridges = [p for p in biblio if p.text.startswith("Bridges, D.")][0]
bogacz = [p for p in biblio if p.text.startswith("Bogacz, R.")][0]
if not (aron._p.getnext() is bogacz._p):
    bogacz._p.getparent().remove(bogacz._p)
    aron._p.addnext(bogacz._p)
    print("moved Bogacz after Aron")

biblio = [p for p in doc.paragraphs
          if p.style.name.startswith("Bibliography") and p.text.strip()]
liu = [p for p in biblio if p.text.startswith("Liu, Z.")][0]
markus = [p for p in biblio if p.text.startswith("Markus, H. R.")][0]
if not (liu._p.getnext() is markus._p):
    markus._p.getparent().remove(markus._p)
    liu._p.addnext(markus._p)
    print("moved Markus & Kitayama after Liu")

doc.save(DST)

# ---------------------------------------------------------------- verify
doc = docx.Document(DST)
refs = [p.text.strip() for p in doc.paragraphs
        if p.style.name.startswith("Bibliography") and p.text.strip()]
print("\nreference list now (%d entries):" % len(refs))
for r in refs:
    print("  -", r[:100])
